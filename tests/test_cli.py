"""Stage 8: config, scto-pull (HTTP mocked) and scto-build end to end."""

import io
import json
import zipfile

import pyarrow.parquet as pq
import pytest
import responses
from farm_form import ROWS, SENTINELS, XML
from openpyxl import load_workbook

from surveycto_python_dict.api import Client, MissingEnvVar
from surveycto_python_dict.cli import build, env_sentinels, load_config, pull
from surveycto_python_dict.codebook import REVIEW_COLUMNS

BASE = "https://acme.surveycto.com"
_ENCRYPTED_XML = XML.replace(
    "<model>", '<model><submission base64RsaPublicKey="MIIBfake"/>'
)
_FORM_LIST = b"""<?xml version="1.0"?>
<xforms xmlns="http://openrosa.org/xforms/xformsList">
  <xform><formID>farm</formID><version>3</version><datasets></datasets>
    <files>{"files":[{"filename":"prev_round.csv","timestamp":1}]}</files></xform>
</xforms>"""
_MANIFEST = b"""<?xml version="1.0"?>
<manifest xmlns="http://openrosa.org/xforms/xformsManifest">
  <mediaFile><filename>prev_round.zip</filename><hash>md5:x</hash>
    <downloadUrl>https://acme.surveycto.com/forms/farm/files/prev_round.zip</downloadUrl></mediaFile>
</manifest>"""


def _config(tmp_path, extra=""):
    path = tmp_path / "config.toml"
    path.write_text(f'form_id = "farm"\nout_dir = "out"\n{extra}', encoding="utf-8")
    return path


# -- config ------------------------------------------------------------------------


def test_load_config_defaults_and_relative_out_dir(tmp_path):
    cfg = load_config(_config(tmp_path))
    assert cfg.form_id == "farm"
    assert cfg.out_dir == tmp_path / "out"
    assert cfg.pii.score_threshold == 0.4


def test_load_config_pii(tmp_path):
    cfg = load_config(
        _config(
            tmp_path,
            '[pii]\nwhitelist = ["owner"]\nmax_distinct = 50\n'
            'languages = { es = "es_core_news_lg" }\n',
        )
    )
    assert cfg.pii.whitelist == ("owner",)
    assert cfg.pii.max_distinct == 50
    assert cfg.pii.languages == {"es": "es_core_news_lg"}


def test_env_sentinels(monkeypatch):
    monkeypatch.setenv("SCTO_SENTINELS", "-99=Don't know; -88=Refused;-1=Missing;")
    assert env_sentinels() == {-99: "Don't know", -88: "Refused", -1: "Missing"}


def test_env_sentinels_missing_raises_naming_the_variable(monkeypatch):
    monkeypatch.delenv("SCTO_SENTINELS", raising=False)
    with pytest.raises(MissingEnvVar, match="SCTO_SENTINELS"):
        env_sentinels()


@pytest.mark.parametrize("raw", ["-99", "-99=", "dk=Don't know"])
def test_env_sentinels_malformed_raises(monkeypatch, raw):
    monkeypatch.setenv("SCTO_SENTINELS", raw)
    with pytest.raises(ValueError):
        env_sentinels()


def test_config_sentinels_table_is_rejected(tmp_path):
    with pytest.raises(ValueError, match="sentinels"):
        load_config(_config(tmp_path, '[sentinels]\n"-99" = "x"\n'))


@pytest.mark.parametrize("extra", ["typo_key = 1\n", '[pii]\nwhitelst = ["x"]\n'])
def test_load_config_unknown_key_raises(tmp_path, extra):
    with pytest.raises(ValueError, match="typo_key|whitelst"):
        load_config(_config(tmp_path, extra))


# -- pull --------------------------------------------------------------------------


def _mock_server(xml: str, wide_method: str):
    responses.get(f"{BASE}/formList", body=_FORM_LIST)
    responses.get(f"{BASE}/forms/farm/xml", body=xml.encode())
    responses.get(f"{BASE}/forms/farm/manifest", body=_MANIFEST)
    zipped = io.BytesIO()
    with zipfile.ZipFile(zipped, "w") as z:
        z.writestr("prev_round.csv", b"phone,owner\n")
    responses.get(f"{BASE}/forms/farm/files/prev_round.zip", body=zipped.getvalue())
    responses.add(wide_method, f"{BASE}/api/v2/forms/data/wide/json/farm", json=ROWS)


@responses.activate
def test_pull_encrypted_form_uses_the_key(tmp_path, monkeypatch):
    key = tmp_path / "k.pem"
    key.write_bytes(b"KEY")
    monkeypatch.setenv("SCTO_KEY_PATH", str(key))
    _mock_server(_ENCRYPTED_XML, responses.POST)
    raw = pull(load_config(_config(tmp_path)), Client("acme", "u", "p"))
    assert raw == tmp_path / "out" / "farm" / "raw"
    assert (raw / "form.xml").read_text(encoding="utf-8") == _ENCRYPTED_XML
    assert json.loads((raw / "wide.json").read_text(encoding="utf-8")) == ROWS
    assert (raw / "datasets" / "prev_round.csv").read_bytes() == b"phone,owner\n"
    assert (raw / "formList.xml").exists()


@responses.activate
def test_pull_unencrypted_form_needs_no_key(tmp_path, monkeypatch):
    monkeypatch.delenv("SCTO_KEY_PATH", raising=False)
    _mock_server(XML, responses.GET)
    raw = pull(load_config(_config(tmp_path)), Client("acme", "u", "p"))
    assert json.loads((raw / "wide.json").read_text(encoding="utf-8")) == ROWS


@responses.activate
def test_pull_encrypted_form_without_key_raises(tmp_path, monkeypatch):
    monkeypatch.delenv("SCTO_KEY_PATH", raising=False)
    _mock_server(_ENCRYPTED_XML, responses.POST)
    with pytest.raises(MissingEnvVar, match="SCTO_KEY_PATH"):
        pull(load_config(_config(tmp_path)), Client("acme", "u", "p"))


# -- build -------------------------------------------------------------------------


@pytest.fixture
def built(tmp_path, analyzer):
    cfg = load_config(_config(tmp_path))
    raw = cfg.out_dir / "farm" / "raw"
    raw.mkdir(parents=True)
    (raw / "form.xml").write_text(XML, encoding="utf-8")
    (raw / "wide.json").write_text(json.dumps(ROWS), encoding="utf-8")
    return build(cfg, SENTINELS, analyzer), cfg.out_dir / "farm"


def test_build_writes_every_output(built):
    outputs, form_dir = built
    assert {k: p.relative_to(form_dir).as_posix() for k, p in outputs.items()} == {
        "questions": "farm_questions.json",
        "sections": "sections",
        "parquet": "farm.parquet",
        "dictionary": "farm_variable_dictionary.json",
        "graph": "farm_variable_graph.json",
        "pii_review": "farm_pii_review.xlsx",
        "dataset_codebook": "farm_dataset_codebook.json",
    }
    assert all(p.exists() for p in outputs.values())
    assert (form_dir / "sections" / "section_plot.json").exists()


def test_build_outputs_agree(built):
    outputs, _ = built
    table = pq.read_table(outputs["parquet"])
    vardict = json.loads(outputs["dictionary"].read_text(encoding="utf-8"))
    assert table.num_rows == len(ROWS)
    assert list(vardict["variables"]) == table.schema.names
    assert all("pii" in v for v in vardict["variables"].values())
    flagged = [
        n
        for n, v in vardict["variables"].items()
        if v["pii"]["status"] in ("contains_pii", "may_contain_pii")
    ]
    review = load_workbook(outputs["pii_review"])
    assert review.sheetnames == ["Columns", "Values"]
    columns = list(review["Columns"].iter_rows(values_only=True))
    assert columns[0] == tuple(REVIEW_COLUMNS)
    assert [row[0] for row in columns[1:]] == flagged
    values = list(review["Values"].iter_rows(values_only=True))
    assert values[0] == ("KEY", *flagged)
    assert len(values) == 1 + len(ROWS)
    assert [row[0] for row in values[1:]] == table.column("KEY").to_pylist()


def test_build_dictionary_holds_no_flagged_value(built):
    outputs, _ = built
    text = outputs["dictionary"].read_text(encoding="utf-8")
    for value in ("Ana Maria Gomez Restrepo", "3104567890", "Jorge Perez"):
        assert value not in text
