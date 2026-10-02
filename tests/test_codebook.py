"""PI-facing codebook: Variables and Summary stats sheets."""

import json

import openpyxl
import pytest
from farm_form import ROWS, SENTINELS, XML

from surveycto_python_dict.cli import build, load_config
from surveycto_python_dict.codebook import plain_label, write_codebook


@pytest.fixture
def workbook(tmp_path, analyzer):
    config = tmp_path / "config.toml"
    config.write_text('form_id = "farm"\nout_dir = "out"\n', encoding="utf-8")
    cfg = load_config(config)
    cfg.raw_dir.mkdir(parents=True)
    (cfg.raw_dir / "form.xml").write_text(XML, encoding="utf-8")
    (cfg.raw_dir / "wide.json").write_text(json.dumps(ROWS), encoding="utf-8")
    build(cfg, SENTINELS, analyzer)
    path = write_codebook(cfg.form_dir, "farm", SENTINELS)
    assert path == cfg.form_dir / "farm_codebook.xlsx"
    return openpyxl.load_workbook(path)


def _rows(sheet):
    header, *rows = [list(r) for r in sheet.iter_rows(values_only=True)]
    return header, {r[0]: dict(zip(header, r, strict=True)) for r in rows}


def test_two_sheets(workbook):
    assert workbook.sheetnames == ["Variables", "Summary stats"]


def test_variables_sheet(workbook):
    header, rows = _rows(workbook["Variables"])
    assert header[:4] == ["variable", "question", "label", "type"]
    assert len(rows) == 29  # every parquet column of the farm form
    area = rows["area"]
    assert (area["label"], area["type"], area["relevance"]) == (
        "Area",
        "decimal",
        "${plots} > 0",
    )
    assert rows["tools_77"]["value_labels"] == "1 = Other; 0 = not chosen"
    assert rows["crop"]["value_labels"] == "01 = Maize; 02 = Rice"
    assert rows["plots"]["sentinels"] == "-99 Don't know: 1"
    assert rows["owner"]["pii_status"] == "may_contain_pii"
    assert rows["KEY"]["type"] == "system"


def test_summary_stats_sheet_keeps_numbers_as_numbers(workbook):
    header, rows = _rows(workbook["Summary stats"])
    assert header[0] == "variable"
    assert rows["plots"]["mean"] == 2
    assert rows["area"]["eligible"] == 1
    assert rows["tools_2"]["selected_share"] == 1.0


def test_sheets_are_filterable_with_frozen_header(workbook):
    for sheet in workbook.worksheets:
        assert sheet.freeze_panes == "A2"
        assert sheet.auto_filter.ref is not None


def test_no_flagged_value_in_any_cell(workbook):
    cells = " ".join(
        str(v)
        for ws in workbook.worksheets
        for row in ws.iter_rows(values_only=True)
        for v in row
    )
    for value in ("Ana Maria Gomez Restrepo", "3104567890", "Jorge Perez", "LT-0042"):
        assert value not in cells


def test_plain_label_strips_html():
    assert (
        plain_label("Nivel de inglés <b>escribiendo</b>")
        == "Nivel de inglés escribiendo"
    )
    assert plain_label("Line one<br/>line two") == "Line one line two"
