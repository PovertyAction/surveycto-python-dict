"""The library path: dictionary, PII rules, stats and codebook from objects, no files."""

import io
import json

import pytest
from farm_form import ROWS, SENTINELS, XML
from openpyxl import load_workbook
from test_stats import _build

from surveycto_python_dict.cast import cast_wide
from surveycto_python_dict.codebook import save_codebook
from surveycto_python_dict.dictionary import build_dictionary
from surveycto_python_dict.pii import PiiConfig, attach, classify, shows_values
from surveycto_python_dict.questions import build_questions
from surveycto_python_dict.stats import stats_rows, summary_stats
from surveycto_python_dict.xml_contract import parse_contract


@pytest.fixture
def built(tmp_path):
    path = tmp_path / "farm.xml"
    path.write_text(XML, encoding="utf-8")
    contract = parse_contract(path)
    questions = build_questions(contract)
    table = cast_wide(ROWS, contract)
    vardict = build_dictionary(table, questions, contract, SENTINELS)
    blocks = classify(table, vardict, questions, PiiConfig(), None)
    attach(vardict, blocks, questions, [])
    return table, questions, vardict, blocks


def test_rules_alone_leave_open_columns_flagged_unscanned(built):
    *_, blocks = built
    assert blocks["owner"] == {
        "status": "may_contain_pii",
        "basis": ["type:text", "unscanned"],
    }
    assert not shows_values(blocks["owner"]["status"])
    # The rules still decide what they can.
    assert blocks["village"]["status"] == "cleared_by_form"
    assert blocks["plots"]["status"] == "cleared_by_form"
    assert blocks["deviceid"]["status"] == "contains_pii"
    assert blocks["visit"]["status"] == "whitelisted"


def test_stats_rows_follow_skip_logic(built):
    table, questions, vardict, _ = built
    stats = {r["variable"]: r for r in stats_rows(table, questions, vardict, SENTINELS)}
    area = stats["area"]
    assert area["relevance"]
    assert area["eligible"] + area["skipped_by_logic"] == area["rows"]


def test_stats_rows_match_summary_stats(tmp_path, analyzer):
    cfg = _build(tmp_path, analyzer)
    form_dir = cfg.form_dir("farm")
    questions = json.loads((form_dir / "farm_questions.json").read_text("utf-8"))
    vardict = json.loads(
        (form_dir / "farm_variable_dictionary.json").read_text("utf-8")
    )
    import pyarrow.parquet as pq

    table = pq.read_table(form_dir / "farm.parquet")
    assert stats_rows(table, questions, vardict, SENTINELS) == summary_stats(
        form_dir, "farm", SENTINELS
    )


def test_save_codebook_to_a_file_object(built):
    table, questions, vardict, _ = built
    buffer = io.BytesIO()
    save_codebook(buffer, vardict, stats_rows(table, questions, vardict, SENTINELS))

    workbook = load_workbook(io.BytesIO(buffer.getvalue()), read_only=True)
    assert workbook.sheetnames == ["Variables", "Summary stats"]


def test_parse_contract_reads_bytes(tmp_path):
    path = tmp_path / "farm.xml"
    path.write_text(XML, encoding="utf-8")
    assert (
        parse_contract(XML.encode()).nodes.keys() == parse_contract(path).nodes.keys()
    )
