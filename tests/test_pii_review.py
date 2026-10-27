"""The review workbook: a person's decisions and reasons, read back on the next build."""

import json

import pytest
from farm_form import SENTINELS, write_pull
from openpyxl import Workbook, load_workbook

from surveycto_python_dict.cli import build, load_config, with_decisions
from surveycto_python_dict.codebook import (
    DATASET_REVIEW_COLUMNS,
    REVIEW_COLUMNS,
    read_pii_decisions,
)

_OPEN = {"contains_pii", "may_contain_pii", "no_hits"}


def _build(tmp_path, analyzer, extra=""):
    path = tmp_path / "config.toml"
    path.write_text(f'forms = ["farm"]\nout_dir = "out"\n{extra}', encoding="utf-8")
    cfg = load_config(path)
    write_pull(cfg)
    build(cfg, SENTINELS, analyzer)
    return cfg


def _review(cfg):
    return cfg.form_dir("farm") / "farm_pii_review.xlsx"


def _rows(cfg, sheet):
    header, *rows = load_workbook(_review(cfg))[sheet].iter_rows(values_only=True)
    return {r[0]: dict(zip(header, r, strict=True)) for r in rows}


def _decide(cfg, sheet, name, decision, reason):
    """Fill in a row as a person would in Excel."""
    workbook = load_workbook(_review(cfg))
    ws = workbook[sheet]
    header = [c.value for c in ws[1]]
    (row,) = [r for r in ws.iter_rows(min_row=2) if r[0].value == name]
    row[header.index("decision")].value = decision
    row[header.index("reason")].value = reason
    workbook.save(_review(cfg))


def _pii(cfg, column):
    path = cfg.form_dir("farm") / "farm_variable_dictionary.json"
    return json.loads(path.read_text(encoding="utf-8"))["variables"][column]["pii"]


def _dataset_pii(cfg, name):
    path = cfg.form_dir("farm") / "farm_dataset_codebook.json"
    rows = json.loads(path.read_text(encoding="utf-8"))
    (row,) = [r for r in rows if f"{r['dataset']}.{r['column']}" == name]
    return row["pii"]


def test_every_column_is_listed_with_the_lists_as_decisions(tmp_path, analyzer):
    cfg = _build(
        tmp_path,
        analyzer,
        '[pii]\nblacklist = ["owner"]\nwhitelist = ["prev_round.phone"]\n',
    )
    workbook = load_workbook(_review(cfg))
    assert workbook.sheetnames == ["Columns", "Datasets", "Values"]
    assert next(workbook["Columns"].iter_rows(values_only=True)) == tuple(
        REVIEW_COLUMNS
    )
    assert next(workbook["Datasets"].iter_rows(values_only=True)) == tuple(
        DATASET_REVIEW_COLUMNS
    )
    columns = _rows(cfg, "Columns")
    parquet = json.loads(
        (cfg.form_dir("farm") / "farm_variable_dictionary.json").read_text(
            encoding="utf-8"
        )
    )["variables"]
    assert set(columns) == set(parquet)
    assert columns["owner"]["decision"] == "personal"
    assert columns["plot_code"]["decision"] is None
    datasets = _rows(cfg, "Datasets")
    assert set(datasets) == {"prev_round.phone", "prev_round.owner"}
    assert datasets["prev_round.phone"]["decision"] == "clear"
    # The decision column is a dropdown on both sheets.
    for sheet, header in [
        ("Columns", REVIEW_COLUMNS),
        ("Datasets", DATASET_REVIEW_COLUMNS),
    ]:
        (validation,) = workbook[sheet].data_validations.dataValidation
        assert validation.formula1 == '"personal,clear"'
        letter = "ABCDEFGHIJ"[header.index("decision")]
        assert str(validation.sqref).startswith(f"{letter}2:")


def test_undecided_open_columns_come_first(tmp_path, analyzer):
    cfg = _build(tmp_path, analyzer, '[pii]\nblacklist = ["owner"]\n')
    order = [
        (row["decision"] is not None, row["status"] not in _OPEN)
        for row in _rows(cfg, "Columns").values()
    ]
    assert order == sorted(order)
    assert order[0] == (False, False)
    assert order[-1][0]


def test_decisions_are_read_back_and_kept(tmp_path, analyzer):
    cfg = _build(tmp_path, analyzer)
    _decide(cfg, "Columns", "plot_code", "personal", "codes the owner's plot")
    _decide(cfg, "Columns", "phone", None, "ask the PI")
    _decide(cfg, "Datasets", "prev_round.owner", "clear", "test names only")
    build(cfg, SENTINELS, analyzer)

    assert _pii(cfg, "plot_code") == {"status": "contains_pii", "basis": ["reviewed"]}
    assert _dataset_pii(cfg, "prev_round.owner") == {
        "status": "whitelisted",
        "basis": ["reviewed"],
    }
    columns, datasets = _rows(cfg, "Columns"), _rows(cfg, "Datasets")
    assert (columns["plot_code"]["decision"], columns["plot_code"]["reason"]) == (
        "personal",
        "codes the owner's plot",
    )
    assert (columns["phone"]["decision"], columns["phone"]["reason"]) == (
        None,
        "ask the PI",
    )
    assert datasets["prev_round.owner"]["decision"] == "clear"

    # Emptying a decision in the workbook takes it back.
    _decide(cfg, "Columns", "plot_code", None, None)
    build(cfg, SENTINELS, analyzer)
    assert _pii(cfg, "plot_code")["basis"] != ["reviewed"]


def test_a_decision_against_the_lists_raises(tmp_path, analyzer):
    cfg = _build(tmp_path, analyzer, '[pii]\nwhitelist = ["plot_code"]\n')
    assert _rows(cfg, "Columns")["plot_code"]["decision"] == "clear"
    _decide(cfg, "Columns", "plot_code", "personal", None)
    with pytest.raises(ValueError, match="`plot_code` is clear in"):
        build(cfg, SENTINELS, analyzer)


def _workbook(path, header, rows):
    """Write a review workbook's Columns sheet by hand."""
    workbook = Workbook()
    workbook.active.title = "Columns"
    workbook.active.append(header)
    for row in rows:
        workbook.active.append(row)
    if "decision" in header:
        workbook.create_sheet("Datasets").append(DATASET_REVIEW_COLUMNS)
    path.parent.mkdir(parents=True)
    workbook.save(path)


def test_two_workbooks_deciding_a_name_both_ways_raise(tmp_path):
    path = tmp_path / "config.toml"
    path.write_text('forms = ["a", "b"]\nout_dir = "out"\n', encoding="utf-8")
    cfg = load_config(path)
    header = ["column", "question", "decision", "reason"]
    _workbook(cfg.form_dir("a") / "a_pii_review.xlsx", header, [["hh", "hh", "clear"]])
    _workbook(
        cfg.form_dir("b") / "b_pii_review.xlsx", header, [["hh", "hh", "personal"]]
    )
    with pytest.raises(ValueError, match="`hh` is clear in a_pii_review.xlsx"):
        with_decisions(cfg)


def test_with_decisions_adds_to_the_lists(tmp_path):
    path = tmp_path / "config.toml"
    path.write_text(
        'forms = ["a"]\nout_dir = "out"\n[pii]\nblacklist = ["kid"]\n', encoding="utf-8"
    )
    cfg = load_config(path)
    _workbook(
        cfg.form_dir("a") / "a_pii_review.xlsx",
        ["column", "question", "decision", "reason"],
        [
            ["kid", "kid", "personal", None],
            ["crop", "crop", "clear", "a crop"],
            ["note", "note", None, "?"],
            # Every binary of a select_multiple clear: the question is clear.
            ["tools_1", "tools", "clear", None],
            ["tools_2", "tools", "clear", None],
            # One instance of a repeat left undecided: the question is not decided.
            ["plot_size_1", "plot_size", "personal", None],
            ["plot_size_2", "plot_size", None, None],
        ],
    )
    cfg, reasons = with_decisions(cfg)
    assert cfg.pii.blacklist == ("kid", "plot_size_1")
    assert cfg.pii.whitelist == ("crop", "tools_1", "tools_2", "tools")
    assert reasons == {"crop": "a crop", "note": "?"}


def test_a_decision_outside_the_dropdown_raises(tmp_path):
    path = tmp_path / "f" / "r.xlsx"
    _workbook(
        path, ["column", "question", "decision", "reason"], [["hh", "hh", "yes", None]]
    )
    with pytest.raises(ValueError, match="Columns row 2: decision 'yes'"):
        read_pii_decisions(path)


def test_a_workbook_from_before_decisions_holds_none(tmp_path):
    path = tmp_path / "f" / "r.xlsx"
    _workbook(path, ["column", "label", "status"], [["hh", "Household", "no_hits"]])
    assert read_pii_decisions(path) == {}
