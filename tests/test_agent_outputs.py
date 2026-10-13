"""End to end: no planted value reaches an output an agent reads.

One synthetic project extends the farm form with fields whose values are canaries:
unique invented strings and numbers. It is built, given stats, a codebook and its
database, and every MCP tool is asked about every question and column. No canary
may appear in a JSON, sections, CSV, xlsx, ``_db.json``/``_db.sql`` or ``flow.*``
file, in a tool's answer, or in what the steps print or log. The blacklist forces
the statuses, so the test does not depend on what Presidio finds. The PII review
workbook holds the flagged values for the team by design, so it is not read.
"""

import json

import pytest
from farm_form import ROWS, SENTINELS, XML
from openpyxl import load_workbook

from surveycto_python_dict import codebook, database
from surveycto_python_dict.cli import build, dataset_csvs, load_config
from surveycto_python_dict.stats import write_form_stats
from surveycto_python_dict.survey import SurveyStore, search_forms, trace_lineage

_INSTANCE = (
    "<ref_number/><birth_date/><contact_code/><buyer/><helpers/><seller/><gps/>"
    "<lat/><market/><nickname/>"
)
_BINDS = """
    <bind nodeset="/farm/ref_number" type="int"/>
    <bind nodeset="/farm/birth_date" type="date"/>
    <bind nodeset="/farm/contact_code" type="string"/>
    <bind nodeset="/farm/buyer" type="select1"/>
    <bind nodeset="/farm/helpers" type="select"/>
    <bind nodeset="/farm/seller" type="select1"/>
    <bind nodeset="/farm/gps" type="geopoint"/>
    <bind nodeset="/farm/lat" type="decimal" calculate="selected-at( /farm/gps , 0)"/>
    <bind nodeset="/farm/market" type="select1"/>
    <bind nodeset="/farm/nickname" type="string"/>
"""
_BODY = """
    <input ref="/farm/ref_number"><label>Reference number</label></input>
    <input ref="/farm/birth_date"><label>Birth date</label></input>
    <input ref="/farm/contact_code"><label>Contact code</label></input>
    <select1 ref="/farm/buyer" appearance="search('buyers', 'matches', 'contact_code',  /farm/contact_code )"><label>Buyer</label>
      <item><label>buyer_label</label><value>buyer_id</value></item>
    </select1>
    <select ref="/farm/helpers" appearance="search('roster')"><label>Helpers</label>
      <item><label>member_name</label><value>member_key</value></item>
    </select>
    <select1 ref="/farm/seller" appearance="search('roster')"><label>Seller</label>
      <item><label>member_name</label><value>member_key</value></item>
    </select1>
    <input ref="/farm/gps"><label>Location</label></input>
    <select1 ref="/farm/market"><label>Market</label>
      <item><label>Velloran Quay</label><value>1</value></item><item><label>Ostrid Fen</label><value>2</value></item>
    </select1>
    <input ref="/farm/nickname"><label>Nickname</label></input>
"""
CANARY_XML = (
    XML.replace("<plot jr:template", _INSTANCE + "\n      <plot jr:template")
    .replace('<bind nodeset="/farm/plot" ', _BINDS + '    <bind nodeset="/farm/plot" ')
    .replace('<group ref="/farm/plot">', _BODY + '    <group ref="/farm/plot">')
)
CANARY_ROWS = [
    ROWS[0]
    | {
        "ref_number": "7340981",
        "birth_date": "Feb 3, 1937",
        "contact_code": "QX7731K",
        "buyer": "B1",
        # Tessaly_Brunwick is no longer in the roster: a code not in the list.
        "helpers": "Zorvane_Quillat Tessaly_Brunwick",
        "seller": "Marrowind_Teel",
        "gps": "4.61238 -74.08127 2500 5",
        "lat": "4.61238",
        "market": "1",
        "nickname": "Brindlewick",
    },
    ROWS[1]
    | {
        "ref_number": "7340982",
        "birth_date": "Nov 21, 1941",
        "contact_code": "QX7732K",
        "buyer": "B2",
        "helpers": "Prellis_Omandu",
        "seller": "Prellis_Omandu",
        "gps": "4.70015 -74.15052 2600 4",
        "lat": "4.70015",
        "market": "2",
        "nickname": "Corvinelle",
    },
]
DATASETS = {
    "prev_round.csv": "phone,owner\n",
    "buyers.csv": "contact_code,buyer_id,buyer_label\n"
    "QX7731K,B1,Stall one\nQX7732K,B2,Stall two\nQX9990K,B3,Stall three\n",
    "roster.csv": "member_key,member_name\n"
    "Zorvane_Quillat,Zorvane Quillat\nPrellis_Omandu,Prellis Omandu\n",
}
FORM_LIST = (
    '<?xml version="1.0"?>\n'
    '<xforms xmlns="http://openrosa.org/xforms/xformsList">\n'
    "  <xform><formID>farm</formID><version>3</version><datasets></datasets>\n"
    "    <files>"
    + json.dumps({"files": [{"filename": f, "timestamp": 1} for f in DATASETS]})
    + "</files></xform>\n</xforms>"
)
CONFIG = """forms = ["farm"]
out_dir = "out"
[pii]
blacklist = ["ref_number", "birth_date", "contact_code", "helpers", "seller", "market"]
whitelist = ["buyer"]
"""

# Values from the submissions and the lists withheld from the outputs.
VALUES = [
    "7340981",
    "7340982",
    "1937-02-03",
    "1941-11-21",
    "Feb 3, 1937",
    "QX7731K",
    "QX7732K",
    "Zorvane Quillat",
    "Prellis Omandu",
    "Marrowind_Teel",
    "4.61238",
    "4.70015",
    "-74.08127",
    "Velloran Quay",
    "Ostrid Fen",
]
# Codes of the withheld select_multiple: the build names its binaries by them.
CODES = ["Zorvane_Quillat", "Prellis_Omandu", "Tessaly_Brunwick"]


def _project(tmp_path, config: str):
    (tmp_path / "config.toml").write_text(config, encoding="utf-8")
    cfg = load_config(tmp_path / "config.toml")
    cfg.raw_dir.mkdir(parents=True)
    (cfg.raw_dir / "formList.xml").write_text(FORM_LIST, encoding="utf-8")
    (cfg.raw_dir / "datasets.json").write_text("[]", encoding="utf-8")
    raw = cfg.form_raw("farm")
    (raw / "files").mkdir(parents=True)
    (raw / "form.xml").write_text(CANARY_XML, encoding="utf-8")
    (raw / "wide.json").write_text(json.dumps(CANARY_ROWS), encoding="utf-8")
    for name, text in DATASETS.items():
        (raw / "files" / name).write_text(text, encoding="utf-8")
    return cfg


def _tool_answers(store: SurveyStore) -> dict[str, str]:
    """Return every MCP tool's answer about every question, column and object."""
    stores = {"farm": store}
    answers = {}
    for name in [*store.questions, *store.variables]:
        answers[f"lookup_variable({name})"] = store.lookup_variable(name, True)
        answers[f"get_choice_list({name})"] = store.get_choice_list(name)
        answers[f"get_gate_chain({name})"] = store.get_gate_chain(name)
        answers[f"neighborhood({name})"] = store.get_variable_neighborhood(name, 3)
        if name in store.questions:
            answers[f"trace_lineage({name})"] = trace_lineage(stores, "farm", name)
    answers["lookup_variables"] = store.lookup_variables(list(store.questions), True)
    groups = {g for q in store.questions.values() for g in q["group_path"]}
    for group in groups:
        answers[f"get_group({group})"] = store.get_group(group)
    answers["get_survey_info"] = store.get_survey_info()
    answers["get_repeat_structure"] = store.get_repeat_structure()
    for repeat in store.graph["repeat_groups"]:
        answers[f"get_repeat_structure({repeat})"] = store.get_repeat_structure(repeat)
    answers["get_datasets"] = store.get_datasets(None, True)
    for dataset in {c["dataset"] for c in store.dataset_columns}:
        answers[f"get_datasets({dataset})"] = store.get_datasets(dataset, True)
    answers["get_project"] = store.get_project()
    answers["get_flow"] = store.get_flow()
    for d in store.flow["datasets"]:
        answers[f"get_flow({d['id']})"] = store.get_flow(d["id"], True)
    answers["get_structure_db"] = store.get_structure_db()
    db = json.loads(
        (store.form_dir.parent / f"{store.project}_db.json").read_text(encoding="utf-8")
    )
    for o in db:
        schema, table = (
            o["name"].split(".", 1) if "." in o["name"] else ("main", o["name"])
        )
        answers[f"get_structure_db({schema})"] = store.get_structure_db(schema)
        answers[f"get_structure_db({o['name']})"] = store.get_structure_db(
            schema, table
        )
    for word in ("buyer", "helpers", "seller", "market", "birth", "location", "code"):
        answers[f"search({word})"] = search_forms(stores, word)
        answers[f"search_questions({word})"] = store.search_questions(word)
    return answers


def _files(out_dir) -> dict[str, str]:
    """Return the text of every output an agent reads, by path."""
    texts = {}
    for path in sorted(out_dir.rglob("*")):
        if "raw" in path.relative_to(out_dir).parts or path.is_dir():
            continue
        if path.name.endswith("_pii_review.xlsx"):
            continue
        if path.suffix in (".json", ".csv", ".md", ".sql"):
            texts[str(path)] = path.read_text(encoding="utf-8-sig")
        elif path.suffix == ".xlsx":
            workbook = load_workbook(path, read_only=True)
            texts[str(path)] = "\n".join(
                str(c)
                for ws in workbook
                for row in ws.iter_rows(values_only=True)
                for c in row
            )
    return texts


def test_no_planted_value_reaches_an_agent_output(tmp_path, analyzer, capsys, caplog):
    cfg = _project(tmp_path, CONFIG)
    build(cfg, SENTINELS, analyzer)
    form_dir = cfg.form_dir("farm")
    codebook.write_codebook(
        form_dir, "farm", write_form_stats(form_dir, "farm", SENTINELS)
    )
    flow = json.loads((cfg.out_dir / "flow.json").read_text(encoding="utf-8"))
    database.write_database(
        cfg.out_dir, cfg.project, list(cfg.forms), dataset_csvs(cfg, flow), {}
    )
    store = SurveyStore(form_dir, "farm")
    # The statuses the canaries rely on, forced by the lists or by type.
    status = {c: v["pii"]["status"] for c, v in store.variables.items()}
    assert status["lat"] == "contains_pii"
    assert store.questions["helpers"]["choices_withheld"] == ["reviewed"]
    assert store.questions["market"]["choices_withheld"] == ["reviewed"]
    assert store.questions["buyer"]["choices_withheld"] is None

    files = _files(cfg.out_dir)
    for kind in ("_questions.json", "_summary_stats.csv", "_codebook.xlsx"):
        assert any(path.endswith(kind) for path in files)
    for kind in (
        "_db.json",
        "_db.sql",
        "flow.json",
        "flow.md",
        "section_ungrouped.json",
    ):
        assert any(path.endswith(kind) for path in files)
    answers = _tool_answers(store)
    printed = capsys.readouterr()
    outputs = {
        **files,
        **answers,
        "stdout": printed.out,
        "stderr": printed.err,
        "log": caplog.text,
    }
    found = {
        where: [v for v in VALUES if v in text]
        for where, text in outputs.items()
        if any(v in text for v in VALUES)
    }
    assert found == {}
    # The binaries' names hold the withheld codes: the tools show them by position.
    named = {
        where: [c for c in CODES if c in text]
        for where, text in answers.items()
        if any(c in text for c in CODES)
    }
    assert named == {}
    assert "helpers_<1>" in answers["lookup_variable(helpers)"]


def test_an_exclude_that_fails_on_a_value_does_not_print_it(tmp_path, analyzer, capsys):
    cfg = _project(
        tmp_path,
        CONFIG + '[submissions]\nexclude = "CAST(nickname AS INTEGER) = 1"\n',
    )
    with pytest.raises(ValueError, match="ConversionException") as raised:
        build(cfg, SENTINELS, analyzer)
    printed = capsys.readouterr()
    for text in (str(raised.value), printed.out, printed.err):
        assert "Brindlewick" not in text
