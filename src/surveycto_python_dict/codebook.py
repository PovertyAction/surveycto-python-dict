"""The xlsx outputs: the PI-facing codebook and the PII review workbook.

``<form>_codebook.xlsx`` has a Variables and a Summary stats sheet, built from the
outputs ``scto-build`` writes; it adds no submission values beyond the summary stats,
which give flagged PII columns counts only. ``scto-codebook`` writes
``<form>_summary_stats.csv`` too, from the same stats. ``<form>_pii_review.xlsx`` is the opposite:
the flagged columns' values, for the people cleaning the data, and where they decide
on each column; the next build reads the decisions back.
"""

import datetime as dt
import json
from collections import defaultdict
from pathlib import Path

import pyarrow as pa
from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from surveycto_python_dict.pii import FLAGGED, PiiConfig, shows_values
from surveycto_python_dict.stats import COLUMNS as STATS_COLUMNS
from surveycto_python_dict.stats import write_form_stats
from surveycto_python_dict.survey import gates, plain_label, rule

VARIABLE_COLUMNS = [
    "variable",
    "question",
    "label",
    "hint",
    "type",
    "group",
    "relevance",
    "value_labels",
    "arrow_type",
    "non_null",
    "missing_rate",
    "sentinels",
    "pii_status",
    "pii_basis",
]
_WRAPPED = {
    "label",
    "reason",
    "hint",
    "relevance",
    "value_labels",
    "frequencies",
    "sentinel_detail",
}
_PERCENT = {"missing_rate", "selected_share"}
# One style object each: openpyxl hashes every cell's style into the workbook.
_HEADER_FONT = Font(bold=True)
_HEADER_FILL = PatternFill("solid", fgColor="DDE5F0")
_WRAP = Alignment(wrap_text=True, vertical="top")


def _variable_rows(variables: dict) -> list[list]:
    rows = []
    for name, v in variables.items():
        survey = v["survey"] or {}
        relevance = rule(gates(survey)) if survey else ""
        rows.append(
            [
                name,
                v["question"] or "",
                plain_label(survey.get("question_text") or ""),
                plain_label(survey.get("hint") or ""),
                survey.get("type") or ("system" if v.get("system") else ""),
                survey.get("group_path", ""),
                relevance,
                "; ".join(
                    f"{c['value']} = {c['label']}" for c in v.get("value_labels", [])
                ),
                v["parquet"]["arrow_type"],
                v["non_null_count"],
                v["missing_rate"],
                "; ".join(
                    f"{code} {s['label']}: {s['count']}"
                    for code, s in v.get("sentinels", {}).items()
                ),
                v["pii"]["status"],
                ", ".join(v["pii"]["basis"]),
            ]
        )
    return rows


def _sheet(
    workbook: Workbook,
    title: str,
    header: list[str],
    rows: list[list],
    dropdown: str | None = None,
) -> None:
    """Add a sheet to a write-only workbook: bold frozen header with a filter, each
    column as wide as its longest value (up to 60), long text wrapped, shares as %,
    and the ``dropdown`` column limited to the ``DECISIONS``.
    """
    rows = [[None if value == "" else value for value in row] for row in rows]
    ws = workbook.create_sheet(title)
    for index, name in enumerate(header):
        longest = max(
            [len(name), *(len(str(r[index])) for r in rows if r[index] is not None)]
        )
        ws.column_dimensions[get_column_letter(index + 1)].width = min(longest + 2, 60)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(header))}{len(rows) + 1}"
    if dropdown is not None and rows:
        letter = get_column_letter(header.index(dropdown) + 1)
        validation = DataValidation(
            type="list", formula1=f'"{",".join(DECISIONS)}"', allow_blank=True
        )
        validation.add(f"{letter}2:{letter}{len(rows) + 1}")
        ws.data_validations.append(validation)
    cells = []
    for name in header:
        cell = WriteOnlyCell(ws, name)
        cell.font, cell.fill = _HEADER_FONT, _HEADER_FILL
        cells.append(cell)
    ws.append(cells)
    styled = [name in _WRAPPED or name in _PERCENT for name in header]
    for row in rows:
        ws.append(
            [
                _cell(ws, name, value) if style else value
                for name, value, style in zip(header, row, styled, strict=True)
            ]
        )


def _cell(ws, name: str, value) -> WriteOnlyCell:
    cell = WriteOnlyCell(ws, value)
    if name in _WRAPPED:
        cell.alignment = _WRAP
    if name in _PERCENT and isinstance(value, (int, float)):
        cell.number_format = "0.0%"
    return cell


REVIEW_COLUMNS = [
    "column",
    "question",
    "label",
    "type",
    "status",
    "basis",
    "hit_rows",
    "non_null",
    "decision",
    "reason",
]
DATASET_REVIEW_COLUMNS = [
    "column",
    "read_by",
    "status",
    "basis",
    "hit_rows",
    "non_null",
    "decision",
    "reason",
]
# A person's decision on a column: personal joins the blacklist, clear the whitelist.
DECISIONS = ("personal", "clear")
# Statuses a person should look at first: flagged, or scanned with no hits.
_OPEN = FLAGGED | {"no_hits"}


VALUE_COLUMNS = ["KEY", "column", "value"]
# A sheet's rows, the header included.
_EXCEL_ROWS = 1_048_576


def _decision(names: set[str], config: PiiConfig) -> str:
    """Return the decision the ``[pii]`` lists hold for any of ``names``."""
    if names & set(config.blacklist):
        return "personal"
    if names & set(config.whitelist):
        return "clear"
    return ""


def _review_order(rows: list[list], header: list[str]) -> list[list]:
    """Undecided flagged or scanned rows first, then the other undecided, then the
    decided, each in their order.
    """
    status, decision = header.index("status"), header.index("decision")
    return sorted(rows, key=lambda r: (r[decision] != "", r[status] not in _OPEN))


def write_pii_review(
    table: pa.Table,
    vardict: dict,
    codebook: list[dict],
    config: PiiConfig,
    reasons: dict[str, str],
    path: Path,
) -> Path:
    """Write ``path``, where the people cleaning the data decide on each column.

    The Columns sheet lists every column and the Datasets sheet every
    ``dataset.column`` the form reads, each with a ``decision`` (``personal`` or
    ``clear``, from the ``[pii]`` lists, which hold the workbook's earlier decisions)
    and its ``reason``; undecided flagged or scanned columns come first. The Values
    sheet has one row per value of a flagged column (``KEY``, column, value), which
    filters by column in Excel.

    Text stays text, so IDs keep their leading zeros. Timestamps are UTC, written
    without the zone, which Excel does not store. More values than a sheet holds
    raise.
    """
    variables = vardict["variables"]
    flagged = [
        c for c in table.schema.names if not shows_values(variables[c]["pii"]["status"])
    ]
    columns = []
    for c, v in variables.items():
        survey = v["survey"] or {}
        columns.append(
            [
                c,
                v["question"] or "",
                plain_label(survey.get("question_text") or ""),
                survey.get("type") or ("system" if v.get("system") else ""),
                v["pii"]["status"],
                ", ".join(v["pii"]["basis"]),
                v["pii"].get("presidio", {}).get("hit_rows", ""),
                v["non_null_count"],
                _decision({c, v["question"]} - {None}, config),
                reasons.get(c, ""),
            ]
        )
    datasets = []
    for row in codebook:
        name = f"{row['dataset']}.{row['column']}"
        datasets.append(
            [
                name,
                ", ".join(row["read_by"]),
                row["pii"]["status"],
                ", ".join(row["pii"]["basis"]),
                row["pii"].get("presidio", {}).get("hit_rows", ""),
                row["non_null"],
                _decision({name}, config),
                reasons.get(name, ""),
            ]
        )
    keys = table.column("KEY").to_pylist()
    values = [
        [key, c, v.replace(tzinfo=None) if isinstance(v, dt.datetime) else v]
        for c in flagged
        for key, v in zip(keys, table.column(c).to_pylist(), strict=True)
        if v is not None
    ]
    if len(values) >= _EXCEL_ROWS:
        raise ValueError(
            f"{path.name}: {len(values)} flagged values, more than a sheet's "
            f"{_EXCEL_ROWS - 1} rows; clear columns in [pii] whitelist and rebuild"
        )
    workbook = Workbook(write_only=True)
    _sheet(
        workbook,
        "Columns",
        REVIEW_COLUMNS,
        _review_order(columns, REVIEW_COLUMNS),
        dropdown="decision",
    )
    _sheet(
        workbook,
        "Datasets",
        DATASET_REVIEW_COLUMNS,
        _review_order(datasets, DATASET_REVIEW_COLUMNS),
        dropdown="decision",
    )
    _sheet(workbook, "Values", VALUE_COLUMNS, values)
    workbook.save(path)
    return path


def read_pii_decisions(path: Path) -> dict[str, tuple[str, str]]:
    """Return ``name -> (decision, reason)`` for each row of a review workbook's
    Columns and Datasets sheets with either filled in.

    A question whose columns all have one decision takes it too, since calculates read
    the question's name: a select_multiple's binaries, a repeat's instances. A workbook
    written before decisions (no ``decision`` column) holds none. A decision other than
    ``personal`` or ``clear`` raises.
    """
    workbook = load_workbook(path, read_only=True)
    sheets = {
        title: list(workbook[title].iter_rows(values_only=True))
        for title in ("Columns", "Datasets")
        if title in workbook.sheetnames
    }
    workbook.close()
    if "decision" not in sheets["Columns"][0]:
        return {}
    decisions = {}
    for title, (header, *rows) in sheets.items():
        name, decision, reason = (
            header.index(c) for c in ("column", "decision", "reason")
        )
        for number, row in enumerate(rows, start=2):
            if row[decision] is None and row[reason] is None:
                continue
            if row[decision] not in (*DECISIONS, None):
                raise ValueError(
                    f"{path.name}, {title} row {number}: decision "
                    f"{row[decision]!r} is not {' or '.join(DECISIONS)}"
                )
            decisions[row[name]] = (row[decision] or "", str(row[reason] or ""))
    header, *rows = sheets["Columns"]
    question, decision = header.index("question"), header.index("decision")
    by_question = defaultdict(set)
    for row in rows:
        if row[question] is not None:
            by_question[row[question]].add(row[decision])
    for q, made in by_question.items():
        if len(made) == 1 and None not in made and q not in decisions:
            decisions[q] = (made.pop(), "")
    return decisions


def write_codebook(form_dir: Path, form_id: str, stats: list[dict]) -> Path:
    """Write ``<form>_codebook.xlsx`` next to the outputs from the form's stats rows."""
    form_dir = Path(form_dir)
    dictionary = json.loads(
        (form_dir / f"{form_id}_variable_dictionary.json").read_text(encoding="utf-8")
    )
    workbook = Workbook(write_only=True)
    _sheet(
        workbook, "Variables", VARIABLE_COLUMNS, _variable_rows(dictionary["variables"])
    )
    _sheet(
        workbook,
        "Summary stats",
        STATS_COLUMNS,
        [[row[c] for c in STATS_COLUMNS] for row in stats],
    )
    path = form_dir / f"{form_id}_codebook.xlsx"
    workbook.save(path)
    return path


def main(argv: list[str] | None = None) -> None:
    """Entry point for ``scto-codebook``."""
    from dotenv import load_dotenv

    from surveycto_python_dict.cli import config_arg, env_sentinels

    cfg = config_arg(argv)
    load_dotenv()
    sentinels = env_sentinels()
    for form in cfg.forms:
        stats = write_form_stats(cfg.form_dir(form), form, sentinels)
        path = write_codebook(cfg.form_dir(form), form, stats)
        print(f"[codebook] {form} -> {path}")
