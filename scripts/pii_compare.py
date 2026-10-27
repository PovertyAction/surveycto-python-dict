"""Score the PII classification against a form's labelled columns (#108).

    uv run python scripts/pii_compare.py template <config> <form> <truth.xlsx>
    uv run python scripts/pii_compare.py compare <config> <form> <truth.xlsx> <out.xlsx> [run ...]

``template`` writes one row per column of a built form, structure only, with ``truth``
prefilled from the config's ``[pii]`` lists (blacklist -> yes, whitelist -> no); the
team fills in the rest. ``compare`` classifies the form's raw pull once per run with
empty lists, so the rules and the scan are measured without the team's review, and
writes each column's status next to its label: a miss is a ``yes`` column left
unflagged, a false alarm a ``no`` column flagged. A run is ``default`` (the config's
settings) or options: spaCy models, left-out entities, phone countries and ``lists``
to keep the config's lists, e.g. ``es=es_core_news_sm,en=en_core_web_sm`` or
``xx=xx_ent_wiki_sm,-LOCATION,@CO``. Neither file holds a value from the data.
"""

import json
import sys
import time
from dataclasses import replace
from pathlib import Path

from dotenv import load_dotenv
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.datavalidation import DataValidation

from surveycto_python_dict.cli import classify_form, env_sentinels, load_config
from surveycto_python_dict.pii import FLAGGED, build_analyzer
from surveycto_python_dict.survey import plain_label

TEMPLATE = [
    "column", "truth", "kind", "note", "source", "question", "label", "type",
    "appearance", "group", "calculation", "status", "basis", "non_null",
]  # fmt: skip


def _share(block: dict) -> float | str:
    """Return the largest share of a column's scanned values one entity hit."""
    presidio = block.get("presidio", {})
    if not presidio.get("distinct_scanned"):
        return ""
    hits = max(presidio["entities"].values(), default=0)
    return round(hits / presidio["distinct_scanned"], 3)


def template(config: Path, form: str, out: Path) -> None:
    """Write the truth template for ``form``'s built outputs to ``out``."""
    assert not out.exists(), f"{out} exists"
    cfg = load_config(config)
    form_dir = cfg.form_dir(form)
    variables = json.loads(
        (form_dir / f"{form}_variable_dictionary.json").read_text(encoding="utf-8")
    )["variables"]
    questions = {
        q["variable_name"]: q
        for q in json.loads(
            (form_dir / f"{form}_questions.json").read_text(encoding="utf-8")
        )
    }
    whitelist, blacklist = set(cfg.pii.whitelist), set(cfg.pii.blacklist)

    rows = []
    for column, entry in variables.items():
        names = {column, entry["question"]} - {None}
        black, white = bool(names & blacklist), bool(names & whitelist)
        truth, source = {
            (True, False): ("yes", "blacklist"),
            (False, True): ("no", "whitelist"),
            (True, True): ("", "both lists"),
            (False, False): ("", ""),
        }[black, white]
        survey = entry["survey"] or {}
        q = questions.get(entry["question"]) or {}
        rows.append(
            (
                truth != "",
                entry["variable_order"],
                [
                    column,
                    truth,
                    "",
                    "",
                    source,
                    entry["question"] or "",
                    plain_label(survey.get("question_text") or ""),
                    survey.get("type") or ("system" if entry.get("system") else ""),
                    q.get("appearance") or "",
                    "/".join(survey.get("group_path") or []),
                    survey.get("calculation") or "",
                    entry["pii"]["status"],
                    ", ".join(entry["pii"]["basis"]),
                    entry["non_null_count"],
                ],
            )
        )
    rows.sort(key=lambda r: (r[0], r[1]))

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Columns"
    sheet.append(TEMPLATE)
    for _, _, row in rows:
        sheet.append(row)
    last = len(rows) + 1
    for cells, options in [
        (f"B2:B{last}", "yes,no"),
        (f"C2:C{last}", "name,phone,id,place,free text,other"),
    ]:
        validation = DataValidation(
            type="list", formula1=f'"{options}"', allow_blank=True
        )
        sheet.add_data_validation(validation)
        validation.add(cells)
    sheet.freeze_panes = "C2"
    sheet.auto_filter.ref = f"A1:N{last}"
    workbook.save(out)
    blank = sum(not filled for filled, _, _ in rows)
    print(f"{out}: {len(rows)} columns, {blank} to fill, {len(rows) - blank} prefilled")


def compare(
    config: Path, form: str, truth_path: Path, out: Path, runs: list[str]
) -> None:
    """Classify ``form`` once per run and write the outcomes to ``out``."""
    cfg = load_config(config)
    sentinels = env_sentinels()
    sheet = load_workbook(truth_path, read_only=True)["Columns"]
    header, *rows = sheet.iter_rows(values_only=True)
    truth = {r[header.index("column")]: r[header.index("truth")] for r in rows}
    assert set(truth.values()) <= {"yes", "no"}, "every column needs truth yes or no"

    results, codebooks, totals = {}, {}, []
    for run in runs:
        # Options: lang=model, -ENTITY, @COUNTRY (phone regions), lists (keep the
        # config's whitelist and blacklist).
        options = [] if run == "default" else run.split(",")
        languages = dict(o.split("=") for o in options if "=" in o)
        dropped = {o[1:] for o in options if o.startswith("-")}
        keep = "lists" in options
        pii = replace(
            cfg.pii,
            whitelist=cfg.pii.whitelist if keep else (),
            blacklist=cfg.pii.blacklist if keep else (),
            languages=languages or cfg.pii.languages,
            entities=tuple(e for e in cfg.pii.entities if e not in dropped),
            countries=tuple(o[1:] for o in options if o.startswith("@"))
            or cfg.pii.countries,
        )
        start = time.perf_counter()
        analyzer = build_analyzer(pii)
        built = time.perf_counter()
        _, _, table, _, blocks, codebook = classify_form(
            replace(cfg, pii=pii), form, sentinels, analyzer, {}
        )
        done = time.perf_counter()
        unlabelled = sorted(set(table.schema.names) - set(truth))
        assert not unlabelled, f"columns without a truth label: {unlabelled}"
        results[run] = blocks
        codebooks[run] = codebook
        flagged = {c for c, b in blocks.items() if b["status"] in FLAGGED}
        misses = [c for c in blocks if truth[c] == "yes" and c not in flagged]
        alarms = [c for c in blocks if truth[c] == "no" and c in flagged]
        totals.append(
            {
                "run": run,
                "columns": len(blocks),
                "truth_yes": sum(truth[c] == "yes" for c in blocks),
                "flagged": len(flagged),
                "misses": len(misses),
                "false_alarms": len(alarms),
                "false_hit_rows": sum(
                    blocks[c].get("presidio", {}).get("hit_rows", 0) for c in alarms
                ),
                "analyzer_seconds": round(built - start, 1),
                "classify_seconds": round(done - built, 1),
            }
        )
        print(totals[-1])

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Totals"
    sheet.append(list(totals[0]))
    for row in totals:
        sheet.append(list(row.values()))
    columns = workbook.create_sheet("Columns")
    columns.append(
        ["column", "truth"]
        + [
            f"{d}_{k}"
            for d in runs
            for k in ("outcome", "status", "basis", "hit_rows", "hit_share")
        ]
    )
    for column in results[runs[0]]:
        row = [column, truth[column]]
        for run in runs:
            block = results[run][column]
            hit = block["status"] in FLAGGED
            outcome = {
                ("yes", True): "flagged",
                ("yes", False): "miss",
                ("no", True): "false alarm",
                ("no", False): "clear",
            }[truth[column], hit]
            presidio = block.get("presidio", {})
            row += [
                outcome,
                block["status"],
                ", ".join(block["basis"]),
                presidio.get("hit_rows", ""),
                _share(block),
            ]
        columns.append(row)
    datasets = workbook.create_sheet("Datasets")
    datasets.append(
        ["dataset", "column", "read_by"]
        + [
            f"{d}_{k}"
            for d in runs
            for k in ("status", "basis", "hit_rows", "hit_share")
        ]
    )
    for i, row in enumerate(codebooks[runs[0]]):
        line = [row["dataset"], row["column"], ", ".join(row["read_by"])]
        for run in runs:
            pii = codebooks[run][i]["pii"]
            line += [
                pii["status"],
                ", ".join(pii["basis"]),
                pii.get("presidio", {}).get("hit_rows", ""),
                _share(pii),
            ]
        datasets.append(line)
    workbook.save(out)
    print(f"-> {out}")


if __name__ == "__main__":
    load_dotenv()
    command, *args = sys.argv[1:]
    if command == "template":
        template(Path(args[0]), args[1], Path(args[2]))
    elif command == "compare":
        compare(
            Path(args[0]),
            args[1],
            Path(args[2]),
            Path(args[3]),
            args[4:] or ["default"],
        )
    else:
        raise SystemExit(__doc__)
